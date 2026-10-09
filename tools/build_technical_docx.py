from pathlib import Path
from copy import deepcopy
import ast,zipfile,hashlib,re,json
from lxml import etree as E
from PIL import Image,ImageDraw,ImageFont
ROOT=Path(__file__).resolve().parents[2]
REF=ROOT/'official_reference/lab_ex4_tf/设计参考例程文档.docx'
SRC=ROOT/'anshibo_event_project/tools/build_technical_report.py'
OUT=ROOT/'submission/technical_document/安视播技术报告.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS={'w':W}
def w(tag):return '{'+W+'}'+tag
def node(tag,**attrs):
 x=E.Element(w(tag))
 for k,v in attrs.items():x.set(w(k),str(v))
 return x
with zipfile.ZipFile(REF) as z: parts={n:z.read(n) for n in z.namelist()}
doc=E.fromstring(parts['word/document.xml']);body=doc.find(w('body'));paras=body.findall(w('p'));sect=deepcopy(body.find(w('sectPr')))
plain=lambda p:''.join(p.itertext())
def text(p,t):
 # Keep paragraph and first-run properties, replace only textual content.
 rp=p.find('./w:r/w:rPr',NS);rp=deepcopy(rp) if rp is not None else None
 for c in list(p):
  if c.tag!=w('pPr'):p.remove(c)
 r=node('r')
 if rp is not None:r.append(rp)
 el=node('t');el.text=t;el.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(el);p.append(r)
 return p
samples={'body':paras[36],'h1':paras[35],'h2':paras[43],'h3':paras[75],'caption':paras[38],'toc1':paras[11],'toc2':paras[13]}
def paragraph(t,kind='body'):
 p=text(deepcopy(samples[kind]),t)
 for n in p.xpath('.//w:bookmarkStart|.//w:bookmarkEnd',namespaces=NS):n.getparent().remove(n)
 return p
# Extract the current technical content without executing PDF generation.
blocks=[];started=False
for stmt in ast.parse(SRC.read_text(encoding='utf-8')).body:
 if not isinstance(stmt,ast.Expr) or not isinstance(stmt.value,ast.Call):continue
 c=stmt.value
 if not isinstance(c.func,ast.Name):continue
 fn=c.func.id
 if fn not in ('page','para','heading','table','diagram'):continue
 args=[ast.literal_eval(a) for a in c.args]
 if fn=='page' and args and args[0]=='1 概述':started=True
 if started:blocks.append((fn,args))
# Keep the original cover paragraphs and original page break.
cover=[deepcopy(x) for x in paras[:10]]
for i,t in {1:'安视播',2:'基于安路FPGA的',3:'校园电子广告牌',4:'与应急信息发布系统',6:'开发板型号：HX4S20C',7:'主芯片（FPGA）： EG4S20BG256',8:'2026-10'}.items():text(cover[i],t)
for c in list(body):body.remove(c)
for c in cover:body.append(c)
toc_title=text(deepcopy(paras[10]),'目录');pp=toc_title.find(w('pPr'));pp.append(node('outlineLvl',val=9));body.append(toc_title)
entries=[]
for fn,args in blocks:
 if fn=='page':entries.append((args[0],0 if re.match(r'^\d+ ',args[0]) else 1))
 elif fn=='heading' and re.match(r'^\d+\.\d+',args[0]):entries.append((args[0],1))
# Preserve the template's dot-leader TOC paragraph styles; make page references updateable.
for i,(title,level) in enumerate(entries):
 p=paragraph(title,'toc1' if level==0 else 'toc2')
 if i==0:
  r=node('r');r.append(node('fldChar',fldCharType='begin'));p.insert(1,r)
  r=node('r');e=node('instrText');e.text=' TOC \\o "1-2" \\h \\z \\u ';r.append(e);p.insert(2,r)
  r=node('r');r.append(node('fldChar',fldCharType='separate'));p.insert(3,r)
 r=node('r');r.append(node('tab'));p.append(r)
 r=node('r');r.append(node('fldChar',fldCharType='begin'));p.append(r)
 r=node('r');e=node('instrText');e.text=f' PAGEREF AnshiboH{i} \\h ';r.append(e);p.append(r)
 r=node('r');r.append(node('fldChar',fldCharType='separate'));p.append(r)
 r=node('r');e=node('t');e.text='—';r.append(e);p.append(r)
 r=node('r');r.append(node('fldChar',fldCharType='end'));p.append(r)
 if i==len(entries)-1:
  r=node('r');r.append(node('fldChar',fldCharType='end'));p.append(r)
 body.append(p)
def pagebreak():
 p=node('p');r=node('r');r.append(node('br',type='page'));p.append(r);body.append(p)
pagebreak()
first=True;heading_idx=0
# Tables use the retained reference table paragraph/cell properties.
ref_table=E.fromstring(parts['word/document.xml']).find('.//'+w('tbl'))
cell_sample=ref_table.find('./w:tr/w:tc',NS)
cell_p=cell_sample.find(w('p'))
cell_props=cell_sample.find(w('tcPr'))
def table(headers,rows,widths):
 t=node('tbl');tp=deepcopy(ref_table.find(w('tblPr')));t.append(tp)
 tw=tp.find(w('tblW'))
 if tw is not None:tw.set(w('type'),'dxa');tw.set(w('w'),'9638')
 # Uniform visible grid, preserving template Table Grid style.
 grid=node('tblGrid');t.append(grid)
 total=sum(widths);cols=[round(v/total*9638) for v in widths]
 for v in cols:grid.append(node('gridCol',w=v))
 for ri,row in enumerate([headers]+rows):
  tr=node('tr');trp=node('trPr');trp.append(node('cantSplit'))
  if ri==0:trp.append(node('tblHeader'))
  tr.append(trp)
  for val,cw in zip(row,cols):
   tc=node('tc');cp=deepcopy(cell_props) if cell_props is not None else node('tcPr')
   tcw=cp.find(w('tcW'))
   if tcw is None:tcw=node('tcW');cp.append(tcw)
   tcw.set(w('w'),str(cw));tcw.set(w('type'),'dxa');tc.append(cp)
   q=text(deepcopy(cell_p),str(val));pp=q.find(w('pPr'))
   if pp is not None:
    ind=pp.find(w('ind'))
    if ind is not None:pp.remove(ind)
   if ri==0:
    r=q.find(w('r'));rp=r.find(w('rPr'))
    if rp is None:rp=node('rPr');r.insert(0,rp)
    rp.append(node('b'))
   tc.append(q);tr.append(tc)
  t.append(tr)
 return t
# A real system diagram replaces the reference's placeholder.
image_path=ROOT/'tmp/docx_review/system.png'
im=Image.new('RGB',(1500,490),'white');dr=ImageDraw.Draw(im);font=ImageFont.truetype('C:/Windows/Fonts/simhei.ttf',27)
for i,(a,b) in enumerate([('TF卡','FAT32资源'),('BMP解析','最近邻缩放'),('SDRAM','双缓冲预读'),('转场与图层','参数及字幕'),('HDMI','视频输出')]):
 x=15+i*300;dr.rectangle((x,30,x+270,155),outline='#647486',width=3);dr.text((x+135,67),a,font=font,anchor='mm',fill='black');dr.text((x+135,113),b,font=font,anchor='mm',fill='black')
 if i<4:dr.line((x+270,93,x+298,93),fill='black',width=3);dr.polygon([(x+298,93),(x+287,85),(x+287,101)],fill='black')
dr.rectangle((90,290,660,405),outline='#647486',width=3);dr.text((375,329),'按键与可选串口事件',font=font,anchor='mm',fill='black');dr.text((375,373),'等级仲裁  CRC校验  帧提交',font=font,anchor='mm',fill='black')
dr.rectangle((815,290,1455,405),outline='#647486',width=3);dr.text((1135,329),'FPGA音源  I2S  HDMI音频',font=font,anchor='mm',fill='black');dr.text((1135,373),'应急帧门控与波形叠加',font=font,anchor='mm',fill='black');dr.line((660,345,815,345),fill='black',width=3);dr.line((375,290,375,225,1040,225,1040,155),fill='black',width=3);im.save(image_path)
from docx import Document
from docx.shared import Cm
scratch=Document();scratch.add_picture(str(image_path),width=Cm(16.5));img_p=deepcopy(scratch.paragraphs[-1]._p)
nsimg={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
for blip in img_p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):blip.set('{'+R+'}embed','rIdAnshiboDiagram')
for fn,args in blocks:
 if fn=='page':
  title=args[0];main=bool(re.match(r'^\d+ ',title))
  if main and not first:pagebreak()
  first=False;p=paragraph(title,'h1' if main else 'h2')
 elif fn=='heading':
  title=args[0];numbered=bool(re.match(r'^\d+\.\d+',title));p=paragraph(title,'h2' if numbered else 'body')
  if not numbered:
   pp=p.find(w('pPr'));pp.append(node('keepNext'))
   r=p.find(w('r'));rp=r.find(w('rPr'))
   if rp is None:rp=node('rPr');r.insert(0,rp)
   rp.append(node('b'))
  else:pass
 elif fn=='para':
  p=paragraph(args[0],'caption' if len(args)>1 and args[1]=='small' else 'body')
 elif fn=='table':body.append(table(*args));continue
 elif fn=='diagram':body.append(deepcopy(img_p));continue
 if fn=='page' or (fn=='heading' and re.match(r'^\d+\.\d+',args[0])):
  b=node('bookmarkStart',id=100+heading_idx,name=f'AnshiboH{heading_idx}');p.insert(1,b);p.append(node('bookmarkEnd',id=100+heading_idx));heading_idx+=1
  pp=p.find(w('pPr'));pp.append(node('keepNext'));pp.append(node('outlineLvl',val=0 if fn=='page' and re.match(r'^\d+ ',args[0]) else 1))
 body.append(p)
body.append(sect)
for field in doc.iter(w('instrText')):
 if field.text:field.text=field.text.replace(chr(92)*2,chr(92))
parts['word/document.xml']=E.tostring(doc,xml_declaration=True,encoding='UTF-8',standalone=True)
# Header image, footer page fields, styles, numbering and theme remain intact.
for name,data in list(parts.items()):
 if re.match(r'word/header\d+\.xml',name):
  xml=E.fromstring(data)
  for t in xml.findall('.//'+w('t')):
   if t.text and '设计参考例程1说明文档' in t.text:t.text=t.text.replace('设计参考例程1说明文档','安视播技术说明文档')
  parts[name]=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
settings=E.fromstring(parts['word/settings.xml']);up=settings.find(w('updateFields'))
if up is None:up=node('updateFields');settings.append(up)
up.set(w('val'),'true');parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
rels=E.fromstring(parts['word/_rels/document.xml.rels']);rr=E.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship');rr.set('Id','rIdAnshiboDiagram');rr.set('Type',R+'/image');rr.set('Target','media/anshibo_system.png');parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/media/anshibo_system.png']=image_path.read_bytes()
cts=E.fromstring(parts['[Content_Types].xml']);ctns='http://schemas.openxmlformats.org/package/2006/content-types'
if not any(x.get('Extension')=='png' for x in cts):x=E.SubElement(cts,'{'+ctns+'}Default');x.set('Extension','png');x.set('ContentType','image/png')
parts['[Content_Types].xml']=E.tostring(cts,xml_declaration=True,encoding='UTF-8',standalone=True)
OUT.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
with zipfile.ZipFile(REF) as a,zipfile.ZipFile(OUT) as b:
 assert a.read('word/styles.xml')==b.read('word/styles.xml')
 assert a.read('word/numbering.xml')==b.read('word/numbering.xml')
 assert a.read('word/footer1.xml')==b.read('word/footer1.xml')
 assert b.testzip() is None
from docx import Document
check=Document(OUT);assert len(check.tables)>=9
assert len(entries)==heading_idx
report={'source':str(REF),'source_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'output':str(OUT),'tables':len(check.tables),'toc_entries':len(entries),'styles_preserved':True,'cover_paragraph_format_preserved':True,'footer_fields_preserved':True,'visual_render_verified':False,'pdf_generated':False}
(ROOT/'tmp/docx_review/check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
